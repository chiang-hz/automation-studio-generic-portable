"""Run an Automation Studio workflow with Python, without Automation Studio server.

The workflow is validated before opening Chrome. Unsupported steps fail closed rather than
being silently skipped. Browser work is always restricted to allowed project domains.
"""
from __future__ import annotations
import argparse
import datetime as dt
import json
import os
import re
import sys
import time
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit
from engine import BrowserEngine

ROOT = Path(__file__).resolve().parent
BACKEND = '__BACKEND__'
SUPPORTED = {
    'navigate', 'newTab', 'switchTab', 'click', 'dblclick', 'fill', 'select', 'upload',
    'press', 'hover', 'check', 'uncheck', 'wait', 'assert', 'manual', 'screenshot',
    'savePagePdf', 'download', 'condition', 'loop', 'extractText', 'extractPattern',
    'captureListSnapshot', 'extractWebData', 'saveDocument', 'waitNewFirst',
    'waitNewListItem', 'script', 'clipboard'
}
# Browser-side JavaScript is executed in the website context, never as Python or Node.js.


def load_json(path):
    if not path.exists(): return {}
    with path.open(encoding='utf-8-sig') as stream: return json.load(stream)


def iter_steps(steps):
    for step in steps:
        if not step.get('enabled', True): continue
        yield step
        for key in ('thenSteps', 'elseSteps', 'steps'):
            yield from iter_steps(step.get(key) or [])


def validate(project, backend):
    errors = []
    browser = project.get('browser') or {}
    if backend not in ('selenium', 'playwright'): errors.append('Unknown Python backend')
    if not project.get('allowedDomains'): errors.append('Project allowedDomains is empty')
    if browser.get('authStateEnabled'):
        verify_url, success_selector = browser.get('authVerifyUrl'), browser.get('authSuccessSelector') or ''
        if not verify_url:
            errors.append('Managed Auth requires authVerifyUrl to identify the CMS website')
        elif urlsplit(verify_url).hostname not in project.get('allowedDomains', []) or len(str(success_selector)) > 400:
            errors.append('Managed Auth verification host must be allowed; optional selector must be at most 400 chars')
    for step in iter_steps(project.get('steps') or []):
        kind = step.get('kind')
        if kind not in SUPPORTED:
            errors.append(f"{step.get('name', step.get('id'))}: unsupported {kind!r} (requires original TypeScript runner)")
        if kind == 'saveDocument' and step.get('documentFormat') not in ('json','xlsx','docx',None):
            errors.append(f"{step.get('name')}: unsupported document format")
        if step.get('periodicRefresh') and (not isinstance(step.get('refreshIntervalSeconds',10),int) or not 5 <= step.get('refreshIntervalSeconds',10) <= 600):
            errors.append(f"{step.get('name')}: refresh interval must be between 5 and 600 seconds")
        if kind == 'script' and re.search(r'(?<![\w.])(require\s*\(|process\.|Buffer\.|__dirname\b|import\s*\(|fs\.|child_process\b)',step.get('script') or str(step.get('value') or '')):
            errors.append(f"{step.get('name')}: Node.js/host script cannot run inside browser JavaScript; rewrite to web DOM script")
    # stealth is a browser setting, not a workflow action. Python uses a normal Chrome session.
    # The original JS stealth plugin is NOT transported and no challenge bypass is attempted.
    return errors


def hostname_allowed(url, domains):
    parsed = urlsplit(url)
    if parsed.scheme not in ('http', 'https') or not parsed.hostname: return False
    host = parsed.hostname.lower().rstrip('.')
    for domain in domains:
        value = str(domain).lower().strip().rstrip('.')
        if value.startswith('*.') and host.endswith(value[1:]) and host != value[2:]: return True
        if host == value: return True
    return False


def require_domain(url, project):
    if not hostname_allowed(url, project.get('allowedDomains', [])):
        raise RuntimeError('Blocked navigation outside project allowedDomains: ' + url)


class Runner:
    def __init__(self, project, engine, parameters):
        self.project, self.engine, self.parameters = project, engine, parameters
        self.variables, self.collections, self.completed, self.downloaded = {}, {}, [], []
        # The click on 檢視 can open a sibling Chrome tab; remember which tabs
        # existed before the click so stale reports are never silently selected.
        self._pending_result_click = None
        self.logs = ROOT / 'logs'
        self.logs.mkdir(parents=True, exist_ok=True)
        self.out = ROOT / 'output'
        self.out.mkdir(parents=True, exist_ok=True)

    def expand(self, text):
        def repl(match):
            name = match.group(1).strip()
            return str(self.variables.get(name, self.parameters.get(name, '')))
        return re.sub(r'{{\s*([^}]+)\s*}}', repl, str(text if text is not None else ''))

    def rules(self, step, key='selectors'):
        return [dict(r, value=self.expand(r.get('value', ''))) for r in step.get(key) or []]

    def locate(self, step, key='selectors', timeout=None):
        return self.engine.locate(self.rules(step, key), step.get('matchMode', 'unique'), step.get('matchIndex', 1), timeout or step.get('timeoutMs', 10000), step.get('frame'), step.get('autoFrameSearch',False))

    def _login_or_challenge_visible(self):
        # Inspect the current document only. Never open a verification tab.
        url = (self.engine.url() or '').lower()
        if re.search(r'/(?:login|sign-?in|signin|sso)(?:[/?#]|$)|/connect/authorize|/cdn-cgi/challenge', url):
            return True
        try:
            if self.engine.backend == 'playwright':
                if self.engine.page.locator('input[type="password"]:visible').count():
                    return True
                if self.engine.page.locator('iframe[src*="turnstile"],iframe[src*="recaptcha"]').count():
                    return True
            else:
                return bool(self.engine.browser.execute_script('''
                    const nodes = Array.from(document.querySelectorAll('input[type="password"],iframe[src*="turnstile"],iframe[src*="recaptcha"]'));
                    return nodes.some(el => !!el.getClientRects().length && getComputedStyle(el).visibility !== 'hidden') ||
                        /verify you are human|checking your browser|驗證您是人類|安全驗證|機器人驗證/i.test(document.title);
                '''))
        except Exception:
            pass
        return False

    @staticmethod
    def _auth_log_url(url):
        # Do not print login callback query strings or authorization tokens.
        parsed = urlsplit(str(url or ''))
        return urlunsplit((parsed.scheme, parsed.netloc, parsed.path, '', '')) or str(url)

    def _auth_tab_urls(self):
        # Test stubs and older standalone engines may not support CDP snapshots.
        snapshot = getattr(self.engine, 'auth_tab_urls', None)
        if callable(snapshot):
            return snapshot()
        return {getattr(self.engine, '_tab_id', lambda: 'current')(): self.engine.url()}

    def _auth_tab_id(self):
        return getattr(self.engine, '_tab_id', lambda: 'current')()

    def _is_authenticated_page_candidate(self, url):
        """URL is a sign-in *candidate*, not proof that site cookies are valid."""
        parsed = urlsplit(str(url or ''))
        expected = urlsplit(self.project.get('browser', {}).get('authVerifyUrl')
                            or self.project['targetUrl'])
        if ((parsed.scheme, parsed.hostname, parsed.port) !=
                (expected.scheme, expected.hostname, expected.port)):
            return False
        if parsed.path.lower().rstrip('/') in ('/404', '/403', '/401', '/error', '/access-denied'):
            return False
        if re.search(r'/(?:login|logout|sign-?in|signin|sso)(?:[/?#]|$)|/connect/authorize|/cdn-cgi/challenge',
                     parsed.path.lower()):
            return False
        return True

    def _peer_login_tab_visible(self):
        """Detect a same-site login tab when the workflow page is an error page."""
        current = self.engine.url()
        if urlsplit(current).path.rstrip('/') not in ('/404', '/403', '/401', '/error'):
            return False
        host = urlsplit(self.project.get('browser', {}).get('authVerifyUrl')
                        or self.project['targetUrl']).hostname
        for tab_id, url in self._auth_tab_urls().items():
            parsed = urlsplit(str(url))
            if tab_id != self._auth_tab_id() and parsed.hostname == host and re.search(
                    r'/(?:login|sign-?in|signin|sso)(?:[/?#]|$)', parsed.path.lower()):
                return True
        return False

    def _manual_auth_on_current_tab(self, resume_target=None):
        # Unlike V3, watch *all* Chrome CDP tabs without changing which tab
        # the user is typing in. The site's login flow may finish in a sibling.
        if getattr(self, '_manual_auth_requested', False):
            raise RuntimeError('AUTH_INVALID: login required again after manual login; check CDP Session/Cookie')
        if self.project.get('browser', {}).get('headless'):
            raise RuntimeError('AUTH_REQUIRED: authentication needs a visible Chrome window')
        self._manual_auth_requested = True
        initial_tab = self._auth_tab_id()
        initial_urls = self._auth_tab_urls()
        print('[AUTH] Login/challenge detected; monitoring all CDP Chrome tabs:',
              ', '.join(self._auth_log_url(u) for u in initial_urls.values()), flush=True)
        print('[AUTH] Complete sign-in in THIS CDP Chrome (including a site-opened new tab). '
              'Workflow auto-resumes within 180s; no Enter required.', flush=True)
        deadline = time.monotonic() + 180
        candidate, stable = None, 0
        last_status = None
        while time.monotonic() < deadline:
            urls = self._auth_tab_urls()
            # Ignore stale pre-existing dashboard tabs. A different tab is only
            # eligible if it was created or navigated after sign-in started.
            eligible = [(tab_id, url) for tab_id, url in urls.items()
                        if self._is_authenticated_page_candidate(url)
                        and (tab_id == initial_tab or initial_urls.get(tab_id) != url)]
            # The original login tab must transition away from its login URL.
            # If it already displayed a protected page at entry, it cannot be
            # used to bypass the login currently happening elsewhere.
            eligible = [(tab_id, url) for tab_id, url in eligible
                        if tab_id != initial_tab or initial_urls.get(tab_id) != url]
            def rank(entry):
                _, url = entry
                path = urlsplit(url).path
                if resume_target and path == urlsplit(resume_target).path: return 3
                if path.startswith('/dataset/provider/'): return 2
                return 1
            eligible.sort(key=rank, reverse=True)
            chosen = eligible[0][0] if eligible else None
            if chosen is not None and chosen == candidate:
                stable += 1
            else:
                candidate, stable = chosen, 1 if chosen is not None else 0
            status = tuple(sorted(self._auth_log_url(u) for u in urls.values()))
            if status != last_status:
                print('[AUTH] CDP tab URLs:', ' | '.join(status) or '(none)', flush=True)
                last_status = status
            if stable >= 3:
                self.engine.select_auth_tab(candidate)
                # A URL transition alone does not prove login; do a final
                # check on the selected tab before restoring the target.
                if self._login_or_challenge_visible():
                    candidate, stable = None, 0
                    time.sleep(.7)
                    continue
                break
            time.sleep(.7)
        else:
            status = ' | '.join(self._auth_log_url(u) for u in self._auth_tab_urls().values())
            raise RuntimeError('AUTH_TIMEOUT: no completed login was detected in the CDP Chrome '
                               f'tabs after 180s. Tabs: {status}. '
                               'Verify that the sign-in finished in this Chrome profile, '
                               'not ordinary Chrome.')
        print('[AUTH] Login completed in Chrome tab:', self._auth_log_url(self.engine.url()), flush=True)
        if resume_target and self.engine.url().rstrip('/') != resume_target.rstrip('/'):
            require_domain(resume_target, self.project)
            print('[AUTH] Resuming original workflow URL:', resume_target, flush=True)
            self.engine.goto(resume_target)  # one recovery navigation only
            if self._login_or_challenge_visible():
                raise RuntimeError(f'AUTH_INVALID: site redirected to login after manual sign-in: '
                                   f'{self._auth_log_url(self.engine.url())}')
        return True

    def ensure_auth(self):
        # Do not navigate to authVerifyUrl, create a __auth_check__ tab, or use
        # a table-header CSS selector as a reason to request sign-in again.
        if not self.project.get('browser', {}).get('authStateEnabled'):
            return
        current = self.engine.url()
        print('[AUTH] Current CDP tab:', current)
        if current == 'about:blank':
            print('[AUTH] This CDP Chrome has a separate login profile; signing in to normal Chrome does not necessarily sign in here.')
        if self._login_or_challenge_visible():
            self._manual_auth_on_current_tab()
        else:
            print('[AUTH] No visible login challenge; workflow will check access on its real target page.')

    def after_action(self):
        require_domain(self.engine.url(), self.project)

    def _capture_result_click(self, step, element):
        """Remember view-link origin without changing, replaying, or re-clicking it."""
        self._pending_result_click = None
        if self.engine.backend != 'selenium':
            return
        try:
            before = self._auth_tab_urls()
            self._pending_result_click = {
                'origin': self._auth_tab_id(),
                'tabs': dict(before),
                'step': str(step.get('name') or ''),
            }
            # Diagnostics do not expose callback parameters or cookie values.
            target = element.get_attribute('target') or '(same tab)'
            href = element.get_attribute('href') or '(script-driven)'
            if urlsplit(str(href)).scheme not in ('http', 'https'):
                href = '(script-driven)'
            print('[TAB] Clicked element; target:', str(target)[:80],
                  'href:', self._auth_log_url(href)[:250], flush=True)
        except Exception as error:
            print('[TAB] Click provenance unavailable:', type(error).__name__, flush=True)

    def _await_report_after_view_click(self, step):
        """Find the report *by its real table*, not by a guessed tab index/URL.

        A Selenium CDP click can leave WebDriver focused on the listing while a
        new tab/window displays the result. We only adopt a tab on the allowed
        CMS host when it contains the exact selector required by this extract.
        Neither login navigation nor a repeated click is performed here.
        """
        rule = step.get('waitBefore') or {}
        if (self.engine.backend != 'selenium' or not self._pending_result_click
                or rule.get('kind') != 'visible' or not rule.get('value')):
            self.wait(rule, step.get('autoFrameSearch', False))
            return
        pending = self._pending_result_click
        self._pending_result_click = None  # never reuse stale popup provenance
        locator = self.expand(rule['value'])
        timeout = int(rule.get('timeoutMs') or step.get('timeoutMs') or 30000)
        deadline = time.monotonic() + timeout / 1000
        origin = pending['origin']
        previous = pending['tabs']
        allowed_hosts = self.project.get('allowedDomains') or []
        last_state = None
        while time.monotonic() < deadline:
            urls = self._auth_tab_urls()
            # A tab opened by an earlier workflow run must not be adopted.
            # Accept the originating tab and newly opened/navigated tabs only.
            candidates = [(handle, url) for handle, url in urls.items()
                          if handle == origin or handle not in previous
                          or previous.get(handle) != url]
            def rank(pair):
                handle, _url = pair
                if handle == origin: return 0
                return 1 if handle not in previous else 2
            candidates.sort(key=rank)
            matches = []
            for handle, url in candidates:
                parsed = urlsplit(str(url))
                if not hostname_allowed(url, allowed_hosts):
                    continue
                if parsed.path.lower().rstrip('/') in ('/login','/404','/403','/401','/error'):
                    continue
                try:
                    self.engine.select_auth_tab(handle)
                    _exists, visible = self.engine.selector_state(
                        [{'strategy': 'css', 'value': locator}],
                        auto_frame=step.get('autoFrameSearch', False))
                    if visible:
                        matches.append((handle, self.engine.url()))
                except Exception:
                    continue  # Tab can close during site-managed navigation.
            if len(matches) == 1:
                selected, url = matches[0]
                self.engine.select_auth_tab(selected)
                print('[TAB] Expected content found in',
                      'current' if selected == origin else 'another',
                      'Chrome tab:', self._auth_log_url(url), flush=True)
                return
            if len(matches) > 1:
                # Do not pick the first of multiple indistinguishable reports.
                targets = ' | '.join(self._auth_log_url(url) for _,url in matches)
                raise RuntimeError('TAB_AMBIGUOUS: multiple tabs contain the '
                                   f'expected content: {targets}')
            # Keep Selenium's active handle aligned with the originating page
            # while waiting, rather than the last inspected unrelated tab.
            try:
                if origin in self.engine.browser.window_handles:
                    self.engine.select_auth_tab(origin)
            except Exception:
                pass
            normalized = tuple(sorted(self._auth_log_url(url) for url in urls.values()))
            if normalized != last_state:
                print('[TAB] Available Chrome tabs:', ' | '.join(normalized) or '(none)', flush=True)
                last_state = normalized
            time.sleep(.3)
        urls = self._auth_tab_urls()
        summary = ' | '.join(self._auth_log_url(url) for url in urls.values()) or '(no tabs)'
        try:
            if origin in self.engine.browser.window_handles:
                self.engine.select_auth_tab(origin)
        except Exception:
            pass
        # Fail closed: no fallback to an unrelated list page or a broad table.
        raise TimeoutError('TAB_RESULT_TIMEOUT: after clicking an element, no allowed Chrome tab '
                           'contained the expected element after '
                           f'{timeout}ms. Required selector: {locator}. '
                           f'Chrome tabs: {summary}. '
                           'Check whether the site opened a popup, a modal, or a different page layout.')

    def wait(self, rule, auto_frame=False, resume_target=None):
        if not rule: return
        kind = rule.get('kind')
        timeout = int(rule.get('timeoutMs') or self.project.get('browser', {}).get('defaultTimeoutMs') or 30000)
        val = self.expand(rule.get('value') or '')
        if kind == 'timeout': time.sleep(timeout / 1000); return
        if kind == 'networkIdle':
            if self.engine.backend == 'playwright': self.engine.page.wait_for_load_state('networkidle', timeout=timeout)
            else:
                deadline=time.monotonic()+timeout/1000
                while time.monotonic()<deadline:
                    if self.engine.browser.execute_script('return document.readyState')=='complete': return
                    time.sleep(.2)
                raise TimeoutError('Document load timed out')
            return
        deadline = time.monotonic() + timeout / 1000
        while time.monotonic() < deadline:
            # A delayed redirect to /login happens while waiting for the
            # workflow element. Pause the wait and resume the target after login.
            if kind in ('visible', 'attached', 'url') and self.project.get('browser', {}).get('authStateEnabled') and (self._login_or_challenge_visible() or self._peer_login_tab_visible()):
                self._manual_auth_on_current_tab(resume_target)
                deadline = time.monotonic() + timeout / 1000
            if kind == 'url' and val in self.engine.url(): return
            if kind in ('visible', 'hidden', 'attached'):
                attached, visible = self.engine.selector_state(
                    [{'strategy': 'css', 'value': val}],
                    auto_frame=auto_frame or rule.get('autoFrameSearch', False))
                if ((kind == 'attached' and attached) or (kind == 'visible' and visible)
                        or (kind == 'hidden' and not visible)):
                    return
            time.sleep(.3)
        raise TimeoutError(f'WAIT_TIMEOUT: {kind} {val}; current URL: {self.engine.url()}')

    def verify(self, rule):
        if not rule: return
        kind = rule.get('kind')
        expected = self.expand(rule.get('expected', ''))
        if kind == 'url':
            if expected not in self.engine.url(): raise RuntimeError('URL verification failed')
            return
        rules = [dict(r, value=self.expand(r.get('value',''))) for r in rule.get('selector') or []]
        if kind in ('hidden', 'exists', 'visible'):
            attached, visible = self.engine.selector_state(rules)
            if kind == 'hidden' and not visible: return
            if kind == 'exists' and attached: return
            if kind == 'visible' and visible: return
            raise RuntimeError('Step verification failed: ' + kind)
        if kind == 'download':
            if self.downloaded: return
            raise RuntimeError('Download verification failed: no artifact saved in this run')
        item = self.engine.locate(rules, timeout=int(rule.get('timeoutMs') or 5000))
        if kind == 'visible' and self.engine.visible(item): return
        if kind == 'exists': return
        if kind == 'text' and expected in self.engine.text(item): return
        if kind == 'value' and expected == self.engine.attr(item, 'value'): return
        if kind == 'checked':
            checked = item.is_checked() if self.engine.backend == 'playwright' else item.is_selected()
            if checked == (str(expected).lower() in ('true','1','yes')): return
        if kind == 'download': return
        raise RuntimeError('Step verification failed: ' + kind)

    def check_condition(self, rule):
        name = rule.get('name') or ''
        source = rule.get('source')
        if source == 'parameter': actual = self.parameters.get(name, '')
        elif source == 'variable': actual = self.variables.get(name, '')
        elif source == 'url': actual = self.engine.url()
        elif source == 'element':
            try: actual = self.engine.text(self.engine.locate([dict(r, value=self.expand(r.get('value',''))) for r in rule.get('selector') or []], timeout=1000))
            except Exception: actual = ''
        else: raise RuntimeError('Unknown condition source')
        op = rule.get('operator')
        target = self.expand(rule.get('value',''))
        if op == 'equals': return str(actual) == str(target)
        if op == 'notEquals': return str(actual) != str(target)
        if op == 'contains': return str(target) in str(actual)
        if op == 'exists': return bool(actual)
        if op == 'notExists': return not bool(actual)
        raise RuntimeError('Unknown condition operator')

    def scan(self, step):
        return self.engine.list_items(self.rules(step), step.get('listLimit',20), step.get('frame'), step.get('autoFrameSearch',False))

    def await_new(self, step, mode):
        source=step.get('sourceVariable') or ''
        baseline=self.variables.get(source,self.parameters.get(source,''))
        if not baseline: raise ValueError('Missing baseline variable for '+mode+': '+source)
        if mode=='waitNewListItem':
            try: previous=json.loads(baseline) if isinstance(baseline,str) else baseline
            except Exception as error: raise ValueError('Baseline is not a JSON snapshot array') from error
            if not isinstance(previous,list): raise ValueError('Baseline must be an array from captureListSnapshot')
            old_fps={str(row.get('fingerprint')) for row in previous if isinstance(row,dict)}
            old_texts={re.sub(r'\s+',' ',str(row.get('text',''))).strip() for row in previous if isinstance(row,dict)}
        else:
            baseline=re.sub(r'\s+',' ',str(baseline)).strip()
        flags=step.get('regexFlags') or 'i'
        supported_flags=set('igmsuyd')
        if any(letter not in supported_flags for letter in flags):raise ValueError('Unsupported regex flags: '+flags)
        pattern=re.compile(step['regexPattern'],re.IGNORECASE if 'i' in flags else 0) if step.get('regexPattern') else None
        timeout=int(step.get('timeoutMs') or (180000 if mode=='waitNewListItem' else 120000))
        deadline=time.monotonic()+timeout/1000
        interval=int(step.get('refreshIntervalSeconds',10)) if step.get('periodicRefresh') else 0
        next_refresh=time.monotonic()+interval if interval else float('inf')
        refresh_count=0
        last=''
        while time.monotonic()<deadline:
            if interval and time.monotonic()>=next_refresh:
                self.engine.reload(min(15000,max(1,int((deadline-time.monotonic())*1000))))
                refresh_count+=1
                next_refresh=time.monotonic()+interval
                self.after_action()
            try:
                if mode=='waitNewFirst':
                    first=self.locate(step, timeout=600)
                    text=re.sub(r'\s+',' ',self.engine.text(first)).strip()
                    last=text
                    if text and text!=baseline and (not pattern or pattern.search(text)):
                        if step.get('outputVariable'): self.variables[step['outputVariable']]=text
                        return text,refresh_count
                else:
                    rows=self.scan(step)
                    last=' | '.join(item['text'] for item in rows[:5])
                    for row in rows:
                        text=row['text']
                        if row['fingerprint'] in old_fps or text in old_texts:continue
                        if pattern and not pattern.search(text):continue
                        if step.get('clickOnMatch'):self.engine.act('click',row['element'])
                        if step.get('outputVariable'):self.variables[step['outputVariable']]=text
                        return text,refresh_count
            except Exception:
                if mode=='waitNewListItem' and step.get('clickOnMatch'):
                    raise
            time.sleep(min(.5,max(.01,deadline-time.monotonic())))
        raise TimeoutError(mode+' timed out. Last visible list: '+last[:500])

    def save_document(self, step):
        import copy
        name=step.get('dataCollection') or 'records'
        rows=self.collections.get(name)
        if rows is None:
            raw=self.variables.get(name,'[]')
            rows=json.loads(raw) if isinstance(raw,str) else raw
        if not isinstance(rows,list) or not rows: raise RuntimeError('Document collection is empty: '+name)
        if len(rows)>50000 or any(not isinstance(r,dict) for r in rows):raise ValueError('Invalid document collection')
        cols=list(dict.fromkeys(k for row in rows for k in row))
        if step.get('documentColumns'):
            requested=[self.expand(col) for col in step['documentColumns']]
            if len(requested)!=len(set(requested)) or any(c not in cols for c in requested): raise ValueError('Document columns are missing or duplicate')
            cols=requested
        fmt=step.get('documentFormat') or 'docx'
        name=self.expand(step.get('documentFileName') or '網頁資料彙整')
        name=re.sub(r'\.(docx|json|xlsx)$','',name,flags=re.I)
        name=re.sub(r'[<>:"/\\|?*\x00-\x1f]','_',name).strip(' .')[:120] or 'data'
        if step.get('documentUseLocalTime') is not False:
            name+='_'+dt.datetime.now().strftime('%Y%m%d_%H%M%S')
        target=self.out/(name+'.'+fmt)
        i=1
        while target.exists(): target=self.out/(name+'_'+str(i)+'.'+fmt);i+=1
        if fmt=='json':target.write_text(json.dumps([{c:r.get(c,'') for c in cols} for r in rows],ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
        elif fmt=='xlsx':
            from openpyxl import Workbook
            wb=Workbook();sheet=wb.active;sheet.title=re.sub(r'[\[\]:*?/\\]','_',self.expand(step.get('documentSheetName') or '資料彙整'))[:31]
            sheet.append(cols)
            for row in rows:sheet.append([str(row.get(k,'')) for k in cols])
            sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
            wb.save(target)
        elif fmt=='docx':
            from docx import Document
            doc=Document();doc.add_heading(self.expand(step.get('documentTitle') or '網頁資料彙整'),0)
            doc.add_paragraph('共 '+str(len(rows))+' 筆資料')
            if step.get('documentLayout')=='paragraphs':
                for ix,row in enumerate(rows):
                    doc.add_paragraph('第 '+str(ix+1)+' 筆',style='Heading 2')
                    for key in cols:doc.add_paragraph(str(key)+'：'+str(row.get(key,'')))
            else:
                table=doc.add_table(rows=1,cols=len(cols));table.style='Table Grid'
                for i,key in enumerate(cols):table.rows[0].cells[i].text=key
                for row in rows:
                    cells=table.add_row().cells
                    for i,key in enumerate(cols):cells[i].text=str(row.get(key,''))
            doc.save(target)
        else:raise RuntimeError('Unsupported document format: '+fmt)
        print('[SAVE]',target)

    def extract_data(self, step):
        fields=step.get('dataFields') or []
        mode=step.get('dataMode') or 'single'
        limit=int(step.get('dataLimit') or 1000)
        if limit<1 or limit>5000:raise ValueError('Data limit must be 1..5000')
        if mode!='table' and not fields:raise ValueError('At least one data field is required')
        selector=self.expand(step.get('dataRowSelector') or '')
        if self.engine.backend=='playwright':
            scope=self.engine.resolve_frame(step.get('frame'),step.get('autoFrameSearch',False),{'strategy':'css','value':selector} if selector else None)
            if mode=='table':
                if not selector:raise ValueError('Table extraction requires dataRowSelector')
                table=scope.locator(selector).first
                table.wait_for(state='attached',timeout=int(step.get('timeoutMs') or 30000))
                cells=table.locator('tr').evaluate_all("""(trs,maxRows)=>trs.slice(0,maxRows+1).map(tr=>[...tr.children]
                    .filter(x=>/^(TD|TH)$/.test(x.tagName)).map(x=>({text:(x.innerText||x.textContent||'').trim(),header:x.tagName==='TH'}))).filter(x=>x.length)""",limit)
                if not cells:raise RuntimeError('No table cells found')
                has_header=any(c['header'] for c in cells[0])
                body=cells[1:] if has_header else cells
                width=max(len(row) for row in cells)
                if width>100:raise RuntimeError('Table exceeds 100 columns')
                columns=[]
                for idx in range(width):
                    base=(cells[0][idx]['text'] if has_header and idx<len(cells[0]) else '') or '欄位'+str(idx+1)
                    key=base; n=2
                    while key in columns or key in ('__proto__','constructor','prototype','_sourceUrl','_capturedAt'):
                        key=base+'_'+str(n); n+=1
                    columns.append(key)
                rows=[{column:str(row[i]['text'] if i<len(row) else '') for i,column in enumerate(columns)} for row in body[:limit]]
            else:
                if mode=='list':
                    if not selector:raise ValueError('List extraction requires dataRowSelector')
                    collection=scope.locator(selector)
                    collection.first.wait_for(state='attached',timeout=int(step.get('timeoutMs') or 30000))
                    roots=[collection.nth(i) for i in range(min(collection.count(),limit))]
                else:roots=[scope]
                rows=[]
                for node in roots:
                    item={}
                    for field in fields:
                        sel=self.expand(field.get('selector') or '')
                        loc=node.locator(sel) if sel else node if mode=='list' else scope.locator('body')
                        count=loc.count()
                        if count>1:raise RuntimeError('Ambiguous web data field: '+field['name'])
                        if count==0:
                            if field.get('required'):raise RuntimeError('Missing required field: '+field['name'])
                            value=''
                        else:
                            el=loc.first
                            kind=field.get('mode') or 'text'
                            value=el.get_attribute(self.expand(field.get('attribute') or '')) if kind=='attribute' else el.inner_html() if kind=='html' else el.inner_text()
                            value=str(value or '').strip()
                            if field.get('required') and not value:raise RuntimeError('Required field empty: '+field['name'])
                        item[field['name']]=value
                    rows.append(item)
            self._store_rows(step,rows)
            return
        # Selenium shares the same Playwright-aware CSS resolver for locating and
        # extracting list/table/field data. The resolver is executed in-page and
        # leaves the original DOM unchanged.
        js="""(opts,css) => {
          const all=(scope,selector)=>css(scope,selector);
          const selected=opts.root?all(document,opts.root):[];
          if(opts.mode==='table') {
            if(!opts.root)throw Error('Table extraction requires a row selector');
            if(!selected.length)throw Error('No table cells found');
            const table=selected[0];
            const cells=all(table,'tr').slice(0,opts.limit+1).map(tr=>[...tr.children]
              .filter(x=>/^(TD|TH)$/.test(x.tagName))
              .map(x=>({text:(x.innerText||x.textContent||'').trim(),header:x.tagName==='TH'})))
              .filter(x=>x.length);
            if(!cells.length)throw Error('No table cells found');
            if(Math.max(...cells.map(x=>x.length))>100)throw Error('Table exceeds 100 columns');
            const hasHeader=cells[0].some(x=>x.header);
            const body=hasHeader?cells.slice(1):cells;
            const cols=[];
            for(let i=0;i<Math.max(...cells.map(x=>x.length));i++) {
              const base=hasHeader?cells[0]?.[i]?.text||'欄位'+(i+1):'欄位'+(i+1);
              let key=base,n=2;
              while(cols.includes(key)||['__proto__','constructor','prototype','_sourceUrl','_capturedAt'].includes(key))key=base+'_'+(n++);
              cols.push(key);
            }
            return body.slice(0,opts.limit).map(c=>Object.fromEntries(cols.map((key,i)=>[key,c[i]?.text||''])));
          }
          if(opts.mode==='list'&&!opts.root)throw Error('List extraction requires a row selector');
          const roots=opts.mode==='list'?selected.slice(0,opts.limit):[document];
          return roots.map(base=>Object.fromEntries(opts.fields.map(field=>{
            const matches=field.selector?all(base,field.selector):opts.mode==='list'?[base]:[document.body];
            if(matches.length>1)throw Error('Ambiguous web data field: '+field.name+' ('+matches.length+' matches)');
            const el=matches[0];
            if(!el){if(field.required)throw Error('Required field missing: '+field.name);return [field.name,''];}
            const value=field.mode==='attribute'?el.getAttribute(field.attribute):field.mode==='html'?el.innerHTML:el.innerText;
            if(field.required&&!String(value||'').trim())throw Error('Required field empty: '+field.name);
            return [field.name,String(value||'').trim()];
          })));
        }"""
        opts={'mode':mode,'root':selector,'limit':limit,'fields':[{**field,'selector':self.expand(field.get('selector') or ''),'attribute':self.expand(field.get('attribute') or '')} for field in fields]}
        # Choose the correct frame by the table/list root, or by the first
        # configured field for a single-record extraction. Querying html would
        # always select the main document even when autoFrameSearch is enabled.
        if step.get('frame') or step.get('autoFrameSearch',False):
            probe = selector or next((f.get('selector') for f in opts['fields'] if f.get('selector')), '')
            if not probe:
                if step.get('autoFrameSearch',False) and not step.get('frame'):
                    raise ValueError('Automatic frame search needs a row or field selector')
                probe = 'html'
            by, probe_value = self.engine._selenium_strategy({'strategy': 'css', 'value': probe})
            matches = self.engine._selenium_find(by, probe_value, step.get('frame'), step.get('autoFrameSearch',False))
            if not matches:
                raise RuntimeError('No matching data source element within the requested frame')
        # Poll dynamic/ExtJS tables rather than failing when the page first renders.
        # Only timeouts caused by missing elements are retried. Invalid selectors
        # and ambiguous fields remain visible errors.
        deadline = time.monotonic() + int(step.get('timeoutMs') or 30000) / 1000
        while True:
            try:
                rows=self.engine.browser.execute_script('const css=(' + self.engine.selenium_css_helper() + ');return ('+js+')(arguments[0],css)',opts)
                if rows:
                    break
            except Exception as error:
                message=str(error)
                if not any(marker in message for marker in ('No table cells found', 'No rows', 'Required field missing', 'Required field empty')):
                    raise
            if time.monotonic() >= deadline:
                raise TimeoutError('Selenium web data extraction timed out: ' + (selector or 'fields'))
            time.sleep(.25)
        self._store_rows(step,rows)

    def _store_rows(self, step, rows):
        if not rows:raise RuntimeError('No web data extracted')
        if step.get('dataIncludeSource') is not False:
            for row in rows:row.update({'_sourceUrl':self.engine.url(),'_capturedAt':dt.datetime.now(dt.timezone.utc).isoformat()})
        collection=step.get('dataCollection') or 'records'
        if not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]{0,79}', collection) or collection in ('__proto__','constructor','prototype'):
            raise ValueError('Invalid data collection name: ' + collection)
        previous=self.collections.get(collection,[]) if step.get('dataAppend') is not False else []
        combined=previous+rows
        if len(combined)>50000:raise RuntimeError('Collection exceeds 50000 rows')
        self.collections[collection]=combined
        self.variables[collection]=json.dumps(combined,ensure_ascii=False)
        self.variables[collection+'_count']=len(combined)
        print('[DATA]',collection,len(rows),'rows')

    def execute(self, step):
        kind = step['kind']
        if kind == 'extractWebData' and self._pending_result_click and self.engine.backend == 'selenium':
            self._await_report_after_view_click(step)
        else:
            # Click provenance applies only to the next extraction; never use an
            # old click to select an unrelated tab several steps later.
            if kind not in ('extractWebData',):
                self._pending_result_click = None
            self.wait(step.get('waitBefore'),step.get('autoFrameSearch',False))
        if kind == 'condition':
            self.run_steps(step.get('thenSteps', []) if self.check_condition(step.get('condition') or {}) else step.get('elseSteps', []))
        elif kind == 'loop':
            values = step.get('loopValues') or []
            if step.get('loopParameter'):
                values = self.parameters.get(step['loopParameter'], '')
                if isinstance(values, str): values = [v.strip() for v in values.split(',') if v.strip()]
            for value in values:
                self.variables[step.get('loopVariable') or 'item'] = value
                self.run_steps(step.get('steps') or [])
        elif kind in ('navigate','newTab'):
            self._pending_result_click = None
            dest=self.expand(step.get('url') or step.get('value') or ('about:blank' if kind=='newTab' else self.project['targetUrl']))
            if dest!='about:blank':require_domain(dest,self.project)
            if kind=='newTab':
                self.engine.new_tab(step.get('tabName'), dest if dest!='about:blank' else None)
            else:
                # Never reload a page already at the intended URL; that can
                # retrigger an identity-provider redirect or Cloudflare.
                if self.engine.url().rstrip('/') == dest.rstrip('/'):
                    print('[AUTH] Skipping navigation: already on the requested URL')
                else:
                    print('[NAV] Opening workflow target:', dest)
                    try:
                        self.engine.goto(dest)
                    except Exception:
                        # Navigation may time out because the site redirected to
                        # a pending login screen. Do not swallow unrelated errors.
                        if not (self.project.get('browser', {}).get('authStateEnabled') and self._login_or_challenge_visible()):
                            raise
                        print('[AUTH] Navigation timed out while site login is showing; waiting for manual sign-in.')
                    print('[NAV] Website landed on:', self.engine.url())
            if self.project.get('browser', {}).get('authStateEnabled') and (self._login_or_challenge_visible() or self._peer_login_tab_visible()):
                self._manual_auth_on_current_tab(dest if kind == 'navigate' else None)
        elif kind == 'switchTab':
            self.engine.switch_tab(step.get('tabTargetMode','index'), self.expand(step.get('tabTarget','')), int(step.get('tabIndex') or 1))
        elif kind in ('click','dblclick','fill','select','upload','press','hover','check','uncheck'):
            value=self.expand(step.get('value',''))
            if step.get('componentPath') or any(r.get('strategy')=='component' for r in step.get('selectors') or []):
                self.engine.component_action(step,value)
            else:
                item=self.locate(step)
                if kind == 'click':
                    self._capture_result_click(step, item)
                else:
                    self._pending_result_click = None
                self.engine.act(kind,item,value)
        elif kind == 'download':
            item=self.locate(step)
            filename=self.expand(step.get('downloadFileName')) if step.get('downloadFileNameMode')=='custom' else None
            timeout=int(step.get('timeoutMs') or self.project.get('browser',{}).get('downloadTimeoutMs',60000))
            mode=step.get('downloadMode') or 'auto'
            expected_format=_download_format_hint(step, filename)
            if mode in ('direct','auto'):
                try: target=self.engine.direct_download(item,filename,timeout,expected_format=expected_format)
                except Exception:
                    if mode=='direct':raise
                    target=self.engine.download(item,filename,timeout,expected_format=expected_format)
            else:target=self.engine.download(item,filename,timeout,expected_format=expected_format)
            self.downloaded.append(target)
            print('[DOWNLOAD]',target)
        elif kind == 'savePagePdf':
            name = self.expand(step.get('pdfFileName') or 'page.pdf')
            name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', '_', name)
            if not name.lower().endswith('.pdf'): name += '.pdf'
            if step.get('pdfUseLocalTime') is not False:
                name = name[:-4] + '_' + dt.datetime.now().strftime('%Y%m%d_%H%M%S') + '.pdf'
            self.engine.pdf(self.out / Path(name).name, step)
        elif kind=='script':
            self.engine.evaluate(self.expand(step.get('script') or step.get('value') or ''))
        elif kind=='clipboard':
            self.engine.clipboard(self.expand(step.get('value') or ''))
        elif kind=='waitNewFirst' or kind=='waitNewListItem':
            text,refreshes=self.await_new(step,kind)
            print('[WAIT NEW]',kind,'matched:',text[:160],'refreshes:',refreshes)
        elif kind == 'screenshot':
            name = self.expand(step.get('value') or 'screenshot.png')
            self.engine.screenshot(self.out / Path(name).name)
        elif kind == 'wait': self.wait(step.get('waitAfter') or {'kind':'timeout','timeoutMs':int(step.get('value') or step.get('timeoutMs') or 1000)})
        elif kind == 'manual':
            if self.project.get('browser',{}).get('headless'):raise RuntimeError('Manual step requires headed Chrome')
            print('[MANUAL]',self.expand(step.get('value') or step.get('description') or step.get('name')))
            mode=step.get('manualCompletionMode') or 'button'
            if mode=='button':input('Complete this action in Chrome, then press Enter...')
            else:
                target=self.expand(step.get('manualExpected') or '')
                if not target:raise ValueError('Manual completion needs manualExpected')
                deadline=time.monotonic()+int(step.get('timeoutMs') or 180000)/1000
                while time.monotonic()<deadline:
                    if mode=='url' and target in self.engine.url():break
                    if mode=='element':
                        try:
                            el=self.engine.locate([{'strategy':'css','value':target}],timeout=700,auto_frame=step.get('autoFrameSearch',False))
                            if self.engine.visible(el):break
                        except Exception:pass
                    time.sleep(.5)
                else:raise TimeoutError('Manual completion condition not met')
        elif kind == 'assert': self.verify(step.get('verification') or {'kind':'visible','selector':step.get('selectors')})
        elif kind == 'extractText':
            item = self.locate(step)
            mode = step.get('extractMode') or 'text'
            self.variables[step.get('outputVariable') or 'text'] = self.engine.attr(item, step.get('extractAttribute') or 'value') if mode == 'attribute' else self.engine.html(item) if mode == 'html' else self.engine.text(item)
        elif kind == 'extractPattern':
            source = str(self.variables.get(step.get('sourceVariable') or '', ''))
            flags = re.IGNORECASE if 'i' in (step.get('regexFlags') or '') else 0
            found = re.search(self.expand(step.get('regexPattern') or ''), source, flags)
            if not found: raise RuntimeError('Pattern not found')
            self.variables[step.get('outputVariable') or 'pattern'] = found.group(int(step.get('regexGroup') or 0))
        elif kind == 'captureListSnapshot':
            rows=self.scan(step)
            if not rows:raise RuntimeError('No visible list rows for baseline snapshot')
            name=step.get('outputVariable')
            if not name:raise ValueError('captureListSnapshot requires outputVariable')
            self.variables[name]=json.dumps([{'fingerprint':r['fingerprint'],'text':r['text']} for r in rows],ensure_ascii=False,separators=(',',':'))
        elif kind == 'extractWebData': self.extract_data(step)
        elif kind == 'saveDocument': self.save_document(step)
        else: raise RuntimeError('Unsupported action: ' + kind)
        if not (kind=='newTab' and self.engine.url()=='about:blank'):
            self.after_action()
        if kind != 'wait': self.wait(step.get('waitAfter'),step.get('autoFrameSearch',False), dest if kind == 'navigate' else None)
        self.verify(step.get('verification'))

    def run_steps(self, steps):
        for step in steps:
            if not step.get('enabled', True): continue
            name = step.get('name') or step.get('id') or 'step'
            print('[STEP]', name)
            started = time.monotonic()
            try:
                self.execute(step)
                self.completed.append({'id':step.get('id'), 'name':name, 'durationSeconds':round(time.monotonic()-started, 2)})
                delay = int(self.project.get('settings',{}).get('minStepDelayMs') or 0)
                if delay: time.sleep(min(delay, 3000)/1000)
            except Exception as exc:
                print('[FAIL]', name, str(exc), file=sys.stderr)
                if self.project.get('settings',{}).get('screenshotMode') != 'never':
                    try: self.engine.screenshot(self.logs / ('failure_' + dt.datetime.now().strftime('%Y%m%d_%H%M%S') + '.png'))
                    except Exception: pass
                # Never blindly retry an action after authentication is lost: it may
                # have already submitted a form or triggered a payment/download.
                raise RuntimeError(f"Workflow stopped at {name}: {exc}. Inspect browser state before replay.") from exc


def _download_format_hint(step, custom_filename=None):
    """Hints for runtime validation; do not assume any .ashx endpoint format."""
    if custom_filename:
        from pathlib import Path
        ext = Path(str(custom_filename)).suffix.lower()
        if ext in {'.xlsx','.xls','.csv','.pdf','.json','.xml','.docx','.zip'}:
            return ext.lstrip('.')
    label = str(step.get('name') or '').lower()
    values = [str(s.get('value') or '').lower() for s in step.get('selectors') or []
              if s.get('strategy') in ('text','role','label')]
    text = ' '.join([label, *values])
    if 'excel' in text or '匯出為excel' in text:
        return 'excel'
    if 'pdf' in text:
        return 'pdf'
    return None



def main(argv=None):
    parser = argparse.ArgumentParser(description='Automation Studio standalone Python runner')
    parser.add_argument('--check-only', action='store_true', help='Validate supported steps without opening browser')
    parser.add_argument('--parameters', default='parameters.json')
    parser.add_argument('--cdp-endpoint', help='Use an already running local CDP Chrome without changing workflow.json')
    args = parser.parse_args(argv)
    project = load_json(ROOT / 'workflow.json')
    if args.cdp_endpoint:
        # Runtime-only override: do not rewrite the exported workflow or secrets.
        from engine import _local_cdp
        endpoint = _local_cdp(args.cdp_endpoint)
        browser = project.setdefault('browser', {})
        browser['connectionMode'] = 'cdp'
        browser['cdpEndpoint'] = endpoint
        browser['headless'] = False
    problems = validate(project, BACKEND)
    if problems:
        for message in problems: print('[UNSUPPORTED]', message, file=sys.stderr)
        return 2
    if args.check_only:
        print('Compatible steps validated for', BACKEND, '(',len(list(iter_steps(project.get('steps') or []))), 'steps )')
        return 0
    params_path = Path(args.parameters)
    if not params_path.is_absolute(): params_path = ROOT / params_path
    supplied = load_json(params_path)
    parameters = {}
    for item in project.get('parameters') or []:
        value = supplied.get(item['name'], item.get('defaultValue',''))
        if item.get('required') and (value is None or value == ''):
            raise ValueError('Missing required parameter: ' + item['name'])
        parameters[item['name']] = value
    engine = BrowserEngine(project, ROOT, BACKEND)
    try:
        engine.open()
        runner = Runner(project, engine, parameters)
        runner.ensure_auth()
        runner.run_steps(project.get('steps') or [])
        print('[DONE] Completed', len(runner.completed), 'steps')
        return 0
    finally:
        engine.close()


if __name__ == '__main__':
    try: sys.exit(main())
    except (KeyboardInterrupt, EOFError):
        print('Execution interrupted or interactive login unavailable', file=sys.stderr)
        sys.exit(130)
    except Exception as exc:
        print('[ERROR]', str(exc), file=sys.stderr)
        sys.exit(1)

