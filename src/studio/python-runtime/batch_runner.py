"""Sequential JSON/CSV/XLSX batch launcher for exported Automation Studio Python workflows.

Uses a single permitted local CDP Chrome, when requested, across consecutive rows.
No passwords, cookies, browser profiles, or unredacted parameter values are written to
result summaries. Transient row parameter files are deleted after each invocation.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import uuid

ROOT = Path(__file__).resolve().parent
ENABLED_COLUMN = '__studio_enabled'
MAX_ROWS = 10000
MAX_INPUT_BYTES = 10 * 1024 * 1024


def as_bool(value, context):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if str(value).strip().lower() in ('true', 'yes', '1', 'on'):
        return True
    if str(value).strip().lower() in ('false', 'no', '0', 'off'):
        return False
    raise ValueError(f'{context}: enabled must be true or false')


def _read_json(file):
    with file.open('r', encoding='utf-8-sig') as stream:
        data = json.load(stream)
    if isinstance(data, list):
        return data
    if not isinstance(data, dict) or not isinstance(data.get('rows'), list):
        raise ValueError('Batch JSON must be an array or an object with a rows array')
    return data['rows']


def _read_csv(file):
    with file.open('r', encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        headers = reader.fieldnames or []
        if not headers or len(set(headers)) != len(headers) or any(not x for x in headers):
            raise ValueError('Batch CSV requires unique, non-empty column names')
        rows = list(reader)
        if any(None in row for row in rows):
            raise ValueError('Batch CSV row has more values than headers')
    return rows


def _read_xlsx(file):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError('XLSX batch files require openpyxl. Run run-batch.bat to install requirements.txt.') from exc
    book = load_workbook(file, read_only=True, data_only=True)
    try:
        rows = book.active.iter_rows(values_only=True)
        headers = [str(x).strip() if x is not None else '' for x in next(rows, ())]
        if not headers or any(not x for x in headers) or len(set(headers)) != len(headers):
            raise ValueError('XLSX first row must contain unique parameter keys')
        result = []
        for values in rows:
            if not values or all(x is None or str(x).strip() == '' for x in values):
                continue
            if len(values) > len(headers) and any(x is not None for x in values[len(headers):]):
                raise ValueError('XLSX row has more values than headers')
            result.append(dict(zip(headers, values)))
            if len(result) > MAX_ROWS:
                raise ValueError(f'Batch list exceeds {MAX_ROWS} rows')
        return result
    finally:
        book.close()


def load_rows(filename, project):
    file = Path(filename).expanduser().resolve()
    if not file.is_file():
        raise FileNotFoundError(f'Batch parameter file not found: {file}')
    if file.stat().st_size > MAX_INPUT_BYTES:
        raise ValueError('Batch parameter file is too large (max 10 MiB)')
    suffix = file.suffix.lower()
    if suffix == '.json':
        raw = _read_json(file)
    elif suffix == '.csv':
        raw = _read_csv(file)
    elif suffix == '.xlsx':
        raw = _read_xlsx(file)
    else:
        raise ValueError('Batch file must end in .json, .csv or .xlsx')
    if len(raw) > MAX_ROWS:
        raise ValueError(f'Batch list exceeds {MAX_ROWS} rows')

    definitions = {param['name']: param for param in project.get('parameters', [])}
    special = {ENABLED_COLUMN, 'enabled', 'id', 'name', 'parameters'}
    result = []
    for index, raw_row in enumerate(raw, start=1):
        if not isinstance(raw_row, dict):
            raise ValueError(f'Row {index}: expected an object')
        enabled = as_bool(raw_row.get('enabled', raw_row.get(ENABLED_COLUMN, True)), f'Row {index}')
        values = raw_row.get('parameters', None)
        if values is None:
            values = {k: v for k, v in raw_row.items() if k not in special}
        if not isinstance(values, dict):
            raise ValueError(f'Row {index}: parameters must be an object')
        unknown = [key for key in values if key not in definitions]
        if unknown:
            raise ValueError(f'Row {index}: unknown parameter keys: {", ".join(unknown[:5])}')
        normalized = {}
        for key, item in definitions.items():
            value = values.get(key, item.get('defaultValue', ''))
            if value is None:
                value = ''
            if item.get('type') == 'boolean' and value != '':
                value = as_bool(value, f'Row {index} parameter {key}')
            elif item.get('type') == 'number' and value != '':
                if not isinstance(value, (int, float)):
                    source = str(value).strip()
                    if not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)', source):
                        raise ValueError(f'Row {index} parameter {key}: invalid number')
                    value = float(source) if '.' in source else int(source)
            elif isinstance(value, (int, float)):
                value = str(value)
            elif not isinstance(value, (str, bool)):
                raise ValueError(f'Row {index} parameter {key}: unsupported data type')
            if enabled and item.get('required') and (value is None or value == ''):
                raise ValueError(f'Row {index}: required parameter {key} is missing')
            if enabled and item.get('type') == 'select' and value != '':
                allowed = item.get('options') or []
                parent = item.get('dependsOn')
                if parent:
                    mapped = (item.get('dependentOptions') or {}).get(str(values.get(parent, normalized.get(parent, ''))))
                    if mapped:
                        allowed = mapped
                if allowed and str(value) not in [str(opt) for opt in allowed]:
                    raise ValueError(f'Row {index} parameter {key}: value is not an allowed option')
            normalized[key] = value
        result.append({'row': index, 'enabled': enabled, 'parameters': normalized})
    return result


def save_json_atomic(file, value):
    file.parent.mkdir(parents=True, exist_ok=True)
    staging = file.with_suffix(file.suffix + '.tmp')
    try:
        staging.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
        os.replace(staging, file)
    finally:
        staging.unlink(missing_ok=True)


def create_run_dir():
    directory = ROOT / '.local' / 'batch-runs' / (dt.datetime.now().strftime('%Y%m%d-%H%M%S') + '-' + uuid.uuid4().hex[:8])
    directory.mkdir(parents=True, exist_ok=False)
    return directory


def _prepare_cdp(project, port=None):
    if sys.platform != 'win32':
        raise RuntimeError('CDP batch launcher requires Windows and an installed Google Chrome')
    import cdp_bootstrap as cdp
    endpoint = cdp.endpoint_for(project, port)
    if cdp.probe(endpoint):
        print('[CDP] Reusing local browser at', endpoint, flush=True)
    else:
        if cdp.tcp_port_open(endpoint):
            raise RuntimeError('CDP port is occupied by another program')
        cdp.wait_for_cdp(endpoint, cdp.start_chrome(cdp.chrome_binary(), endpoint, project))
    if not (project.get('browser') or {}).get('authStateEnabled'):
        print('[LOGIN] Complete website login in this Chrome window before executing batch.', flush=True)
        try:
            input('[LOGIN] Press Enter when ready...')
        except EOFError as exc:
            raise RuntimeError('CDP login requires an interactive terminal') from exc
    return endpoint


def run_row(row, endpoint, folder):
    # Store per-row parameters only in a private, short-lived file, never in logs.
    descriptor, filename = tempfile.mkstemp(prefix='row-', suffix='.json', dir=folder)
    file = Path(filename)
    try:
        with os.fdopen(descriptor, 'w', encoding='utf-8') as stream:
            json.dump(row['parameters'], stream, ensure_ascii=False)
        command = [sys.executable, '-u', str(ROOT / 'runner.py'), '--parameters', str(file)]
        if endpoint:
            command += ['--cdp-endpoint', endpoint]
        env = dict(os.environ, PYTHONUTF8='1')
        log_file = folder / f"row-{row['row']:04d}.log"
        with log_file.open('w', encoding='utf-8') as log:
            process = subprocess.Popen(command, cwd=ROOT, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                       text=True, encoding='utf-8', errors='replace')
            try:
                assert process.stdout is not None
                for line in process.stdout:
                    print(line, end='', flush=True)
                    log.write(line)
                    log.flush()
                code = process.wait()
            except KeyboardInterrupt:
                process.terminate()
                try:
                    process.wait(timeout=5)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait()
                raise
        return code, log_file
    finally:
        file.unlink(missing_ok=True)


def run_batch(argv=None):
    parser = argparse.ArgumentParser(description='Sequential standalone Python batch workflow launcher')
    parser.add_argument('--input', default='batch.json', help='JSON / CSV / XLSX batch parameter file')
    parser.add_argument('--on-error', choices=('stop', 'continue'), default=None, help='Override batch-settings.json (default: continue)')
    parser.add_argument('--cdp', action='store_true', help='Reuse a local Chrome CDP session across rows')
    parser.add_argument('--port', type=int, help='Local CDP port')
    parser.add_argument('--check-only', action='store_true', help='Validate workflow and batch list without starting Chrome')
    parser.add_argument('--dry-run', action='store_true', help='Validate and preview task count without running')
    args = parser.parse_args(argv)
    if args.on_error is None:
        settings_file = ROOT / 'batch-settings.json'
        if settings_file.exists():
            settings = json.loads(settings_file.read_text(encoding='utf-8-sig'))
            strategy = settings.get('onError', 'continue')
            if strategy not in ('continue', 'stop'):
                raise ValueError('Invalid batch-settings.json onError: expected continue or stop')
            args.on_error = strategy
        else:
            args.on_error = 'continue'
    from runner import BACKEND, validate
    project = json.loads((ROOT / 'workflow.json').read_text(encoding='utf-8-sig'))
    errors = validate(project, BACKEND)
    if errors:
        for error in errors:
            print('[UNSUPPORTED]', error, file=sys.stderr)
        return 2
    source = Path(args.input)
    if not source.is_absolute():
        source = ROOT / source
    rows = load_rows(source, project)
    enabled = [row for row in rows if row['enabled']]
    print(f'[BATCH] {len(rows)} rows, {len(enabled)} enabled, {len(rows)-len(enabled)} skipped; mode={"CDP" if args.cdp else "standard"}', flush=True)
    if args.check_only or args.dry_run:
        print('[BATCH] Validation passed. No browser was opened.', flush=True)
        return 0
    if not enabled:
        print('[ERROR] No enabled batch rows; nothing was executed.', file=sys.stderr)
        return 2
    folder = create_run_dir()
    endpoint = _prepare_cdp(project, args.port) if args.cdp else None
    summary = {'mode': 'cdp' if args.cdp else 'standard', 'source': source.name, 'total': len(rows), 'enabled': len(enabled), 'rows': []}
    failed = False
    print('[BATCH] Log folder:', folder, flush=True)
    try:
        for pos, row in enumerate(enabled, 1):
            print(f"[BATCH] Task {pos}/{len(enabled)} (source row {row['row']})", flush=True)
            started = dt.datetime.now(dt.timezone.utc)
            try:
                exit_code, log_file = run_row(row, endpoint, folder)
            except OSError as exc:
                # Starting a child Python process can fail before it writes any log.
                # Record a failed row so --on-error continue remains meaningful.
                exit_code = 1
                log_file = folder / f"row-{row['row']:04d}.log"
                log_file.write_text(f'Unable to start workflow process: {type(exc).__name__}: {exc}\n', encoding='utf-8')
                print('[ERROR] Unable to start the Python workflow for this row.', flush=True)
            success = exit_code == 0
            failed |= not success
            summary['rows'].append({'row': row['row'], 'status': 'success' if success else 'failed',
                                    'exitCode': exit_code, 'startedAt': started.isoformat(),
                                    'finishedAt': dt.datetime.now(dt.timezone.utc).isoformat(), 'log': log_file.name})
            save_json_atomic(folder / 'summary.json', summary)
            print(f"[BATCH] Task {pos}/{len(enabled)} {'OK' if success else 'FAILED'} (exit {exit_code})", flush=True)
            if not success and args.on_error == 'stop':
                print('[BATCH] Stopped on first failed task.', flush=True)
                break
    finally:
        save_json_atomic(folder / 'summary.json', summary)
        with (folder / 'summary.csv').open('w', encoding='utf-8-sig', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=['row', 'status', 'exitCode', 'startedAt', 'finishedAt', 'log'])
            writer.writeheader()
            writer.writerows(summary['rows'])
    print(f"[BATCH] Completed {sum(r['status']=='success' for r in summary['rows'])}/{len(enabled)} enabled task(s). Logs: {folder}", flush=True)
    return 1 if failed else 0


if __name__ == '__main__':
    try:
        sys.exit(run_batch())
    except (ValueError, OSError, RuntimeError, json.JSONDecodeError) as error:
        print('[ERROR]', str(error), file=sys.stderr)
        sys.exit(2)
    except KeyboardInterrupt:
        print('[STOP] Batch interrupted.', file=sys.stderr)
        sys.exit(130)
